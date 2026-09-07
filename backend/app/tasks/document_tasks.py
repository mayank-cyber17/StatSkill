"""
Document processing pipeline — runs in-process using asyncio.
No Redis / Celery required for development or demo.
"""
import asyncio
import os
import re

from app.core.database import AsyncSessionLocal
from app.models.assessment import Document, DocumentChunk
from app.services.ai.rag_service import RAGService
from sqlalchemy import select


# ──────────────────────────────────────────────
# Text extraction helpers
# ──────────────────────────────────────────────

def _extract_pdf(file_path: str) -> str:
    import pymupdf as fitz  # updated API
    text = ""
    with fitz.open(file_path) as doc:
        for page in doc:
            text += page.get_text() + "\n"
    return text


def _extract_docx(file_path: str) -> str:
    from docx import Document as DocxDoc
    doc = DocxDoc(file_path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_pptx(file_path: str) -> str:
    from pptx import Presentation
    prs = Presentation(file_path)
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    parts.append(para.text)
    return "\n".join(parts)


def _extract_text(file_path: str, content_type: str) -> str:
    ct = (content_type or "").lower()
    ext = os.path.splitext(file_path)[1].lower()

    if ct == "application/pdf" or ext == ".pdf":
        return _extract_pdf(file_path)
    elif ct in ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",) or ext == ".docx":
        return _extract_docx(file_path)
    elif ct in ("application/vnd.openxmlformats-officedocument.presentationml.presentation",) or ext == ".pptx":
        return _extract_pptx(file_path)
    elif ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    else:
        # Generic fallback: try plain text
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception:
            return ""


# ──────────────────────────────────────────────
# Chunking helper
# ──────────────────────────────────────────────

def _chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> list[str]:
    """Simple word-based chunking with overlap (no tiktoken required)."""
    words = text.split()
    chunks = []
    step = chunk_size - overlap
    for i in range(0, len(words), step):
        chunk_words = words[i: i + chunk_size]
        chunks.append(" ".join(chunk_words))
        if i + chunk_size >= len(words):
            break
    return chunks


# ──────────────────────────────────────────────
# Main async processing function
# ──────────────────────────────────────────────

async def async_process_document(document_id: int):
    from app.core.config import get_settings
    settings = get_settings()

    async with AsyncSessionLocal() as db:
        result = await db.execute(select(Document).where(Document.id == document_id))
        doc = result.scalar_one_or_none()

        if not doc:
            return

        doc.processing_status = "PROCESSING"
        await db.commit()

        try:
            # Resolve the actual file path
            # file_path is stored as relative sub-path inside UPLOAD_DIR
            candidate = doc.file_path
            if not os.path.isabs(candidate):
                candidate = os.path.join(settings.UPLOAD_DIR, candidate)
            if not os.path.exists(candidate):
                # Try directly
                candidate = doc.file_path

            raw_text = _extract_text(candidate, doc.file_type)

            # Clean whitespace
            raw_text = re.sub(r'\s+', ' ', raw_text).strip()

            if not raw_text:
                raise ValueError("No text could be extracted from the document.")

            # Chunk
            text_chunks = _chunk_text(raw_text, chunk_size=500, overlap=50)

            db_chunks = []
            for idx, chunk_text in enumerate(text_chunks):
                db_chunk = DocumentChunk(
                    document_id=doc.id,
                    chunk_index=idx,
                    content=chunk_text,
                    token_count=len(chunk_text.split()),
                )
                db.add(db_chunk)
                await db.commit()
                await db.refresh(db_chunk)
                db_chunks.append(db_chunk)

            # Add to ChromaDB vector store
            rag_service = RAGService()
            collection_name = f"doc_{doc.id}"

            chroma_chunks = [
                {
                    "id": str(c.id),
                    "content": c.content,
                    "metadata": {"document_id": doc.id, "chunk_index": c.chunk_index},
                }
                for c in db_chunks
            ]
            await rag_service.add_chunks(collection_name, chroma_chunks)

            doc.processing_status = "READY"
            await db.commit()

        except Exception as e:
            print(f"[document_processor] Error processing doc {document_id}: {e}")
            doc.processing_status = "FAILED"
            doc.error_message = str(e)
            await db.commit()


# ──────────────────────────────────────────────
# Entry point called from upload endpoint
# ──────────────────────────────────────────────

def trigger_document_processing(document_id: int):
    """
    Fire-and-forget: schedule the async processor without blocking the caller.
    Works without Redis/Celery.
    """
    asyncio.create_task(async_process_document(document_id))
